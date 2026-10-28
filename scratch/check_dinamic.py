import os
import openpyxl

base_dir = r"C:\Projetos\ARGUS\gestao_projetos\execucao_financeira"
fin_file = os.path.join(base_dir, "execucao_financeira.xlsx")

wb = openpyxl.load_workbook(fin_file, data_only=False)
ws_din = wb["dinamic"]
print("Sheet 'dinamic' max row:", ws_din.max_row, "max col:", ws_din.max_column)
print("Pivot tables:", getattr(ws_din, '_pivots', None))

for r in range(1, ws_din.max_row + 1):
    vals = [ws_din.cell(r, c).value for c in range(1, ws_din.max_column + 1)]
    if any(v is not None for v in vals):
        print(f"R{r}: {vals}")
