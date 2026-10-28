import os
import openpyxl

fin_file = r"C:\Projetos\ARGUS\gestao_projetos\execucao_financeira\execucao_financeira_backup_20261027.xlsx"
test_file = r"C:\Projetos\ARGUS\gestao_projetos\execucao_financeira\execucao_financeira_test.xlsx"

wb = openpyxl.load_workbook(fin_file)
print("Sheetnames:", wb.sheetnames)
ws = wb["extrato_bancario"]
print("Original max row:", ws.max_row)

# Salvar como teste
wb.save(test_file)
print("Salvo com sucesso como test_file:", test_file)

# Recarregar test_file
wb2 = openpyxl.load_workbook(test_file)
print("Recarregado com sucesso! Sheets:", wb2.sheetnames)
ws2 = wb2["dinamic"]
print("Pivot tables in reloaded:", getattr(ws2, '_pivots', None))
