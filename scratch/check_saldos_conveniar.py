import os
import openpyxl

conv_dir = r"C:\Projetos\ARGUS\gestao_projetos\execucao_financeira\dados-conveniar"

for fname in sorted(os.listdir(conv_dir)):
    if not fname.endswith(".xlsx"): continue
    wb = openpyxl.load_workbook(os.path.join(conv_dir, fname), data_only=True)
    ws = wb[wb.sheetnames[0]]
    print("=" * 80)
    print("ARQUIVO:", fname)
    
    # Procurar saldo final do projeto e totais
    for r in range(1, ws.max_row + 1):
        txt = str(ws.cell(r, 2).value or "").strip()
        if "Saldo Final do Projeto" in txt or "Total do Projeto" in txt:
            vals = [ws.cell(r, c).value for c in range(1, 10)]
            print(f"  R{r}: {vals}")
        elif "Rubrica:" in txt:
            saldo_ant = ws.cell(r+1, 4).value
            print(f"  R{r}: {txt} | Saldo Anterior: {saldo_ant}")
