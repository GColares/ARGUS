import os
import openpyxl

conv_dir = r"C:\Projetos\ARGUS\gestao_projetos\execucao_financeira\dados-conveniar"

for fname in sorted(os.listdir(conv_dir)):
    if not fname.endswith(".xlsx"): continue
    print("=" * 80)
    print("ARQUIVO:", fname)
    print("=" * 80)
    wb = openpyxl.load_workbook(os.path.join(conv_dir, fname), data_only=True)
    ws = wb[wb.sheetnames[0]]
    for r in range(1, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, 10)]
        # imprimir linha
        txts = [f"C{c}:{v}" for c, v in enumerate(vals, 1) if v is not None and str(v).strip() != ""]
        if txts:
            print(f"R{r:2d} -> " + " | ".join(txts))
