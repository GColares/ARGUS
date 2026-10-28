import os
import openpyxl

fin_file = r"C:\Projetos\ARGUS\gestao_projetos\execucao_financeira\execucao_financeira_backup_20261027.xlsx"
wb = openpyxl.load_workbook(fin_file, data_only=True)
ws = wb["extrato_bancario"]

for r in range(2, ws.max_row + 1):
    c_val = ws.cell(r, 4).value
    dt_lan = ws.cell(r, 1).value
    dt_pag = ws.cell(r, 2).value
    val = ws_ext.cell(r, 10).value if 'ws_ext' in locals() else ws.cell(r, 10).value
    doc = ws.cell(r, 6).value
    print(f"L{r:2d}: Conta={str(c_val)[:5]} | Lan={dt_lan} | Pag={dt_pag} | Doc={doc} | Val={val}")
