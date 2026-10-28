import os
import openpyxl

base_dir = r"C:\Projetos\ARGUS\gestao_projetos\execucao_financeira"
fin_file = os.path.join(base_dir, "execucao_financeira.xlsx")
conv_dir = os.path.join(base_dir, "dados-conveniar")

wb_fin = openpyxl.load_workbook(fin_file, data_only=True)
ws_ext = wb_fin["extrato_bancario"]

print("=== REGISTROS ATUAIS EM extrato_bancario ===")
rows_by_conta = {}
for r in range(2, ws_ext.max_row + 1):
    dt_lan = ws_ext.cell(r, 1).value
    dt_pag = ws_ext.cell(r, 2).value
    mes = ws_ext.cell(r, 3).value
    conta = ws_ext.cell(r, 4).value
    rub = ws_ext.cell(r, 5).value
    doc = ws_ext.cell(r, 6).value
    mov = ws_ext.cell(r, 7).value
    tipo = ws_ext.cell(r, 8).value
    desc = ws_ext.cell(r, 9).value
    val = ws_ext.cell(r, 10).value
    
    if not any([dt_lan, dt_pag, conta, val]):
        continue
        
    rows_by_conta.setdefault(conta, []).append({
        'row': r, 'dt_lan': dt_lan, 'dt_pag': dt_pag, 'mes': mes,
        'rubrica': rub, 'doc': doc, 'mov': mov, 'tipo': tipo,
        'desc': desc, 'val': val
    })

for c, items in rows_by_conta.items():
    print(f"\nConta '{c}': {len(items)} registros")
    dts = [it['dt_pag'] for it in items if it['dt_pag']]
    if dts:
        print(f"  Período: {min(dts)} até {max(dts)}")
    rubricas = set(it['rubrica'] for it in items)
    print(f"  Rubricas: {rubricas}")
    soma = sum(it['val'] for it in items if it['val'] is not None)
    print(f"  Saldo Total somado: {soma:.2f}")

print("\n" + "=" * 60)
print("VARREDURA COMPLETA DOS ARQUIVOS CONVENIAR")
print("=" * 60)

for fname in sorted(os.listdir(conv_dir)):
    if not fname.endswith(".xlsx"): continue
    fpath = os.path.join(conv_dir, fname)
    print(f"\n>>> ARQUIVO: {fname} <<<")
    wb_c = openpyxl.load_workbook(fpath, data_only=True)
    ws_c = wb_c[wb_c.sheetnames[0]]
    
    current_rubrica = None
    trans_count = 0
    rubrica_sections = []
    
    for r in range(1, ws_c.max_row + 1):
        col_b = str(ws_c.cell(r, 2).value or "").strip()
        
        if col_b.startswith("Rubrica:"):
            current_rubrica = col_b.replace("Rubrica:", "").strip()
            saldo_ant = ws_c.cell(r+1, 4).value if r+1 <= ws_c.max_row else None
            rubrica_sections.append({'rubrica': current_rubrica, 'start_row': r, 'saldo_ant': saldo_ant})
            print(f"  [Rubrica Encontrada] '{current_rubrica}' na linha {r} (Saldo Anterior: {saldo_ant})")
        
        # Verificar se é linha de lançamento (data em col B)
        dt_lan = ws_c.cell(r, 2).value
        # Se for string com data ou datetime
        is_date = False
        if dt_lan:
            if hasattr(dt_lan, 'strftime'):
                is_date = True
            elif any(c in str(dt_lan) for c in ['/', '-']) and not str(dt_lan).startswith("Rubrica") and not str(dt_lan).startswith("Data") and not str(dt_lan).startswith("Saldo"):
                is_date = True
                
        if is_date:
            dt_pag = ws_c.cell(r, 3).value
            doc = ws_c.cell(r, 4).value
            tipo = ws_c.cell(r, 5).value
            hist = ws_c.cell(r, 6).value
            deb = ws_c.cell(r, 7).value
            cred = ws_c.cell(r, 8).value
            saldo = ws_c.cell(r, 9).value
            trans_count += 1
            print(f"    L{r} [{current_rubrica}]: Lan={dt_lan}, Pag={dt_pag}, Doc={doc}, Tipo={tipo}, Deb={deb}, Cred={cred}")
            print(f"       Hist: {str(hist)[:80]}...")
            
    print(f"  Total lançamentos no arquivo: {trans_count}")
