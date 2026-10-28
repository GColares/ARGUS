import os
import openpyxl

base_dir = r"C:\Projetos\ARGUS\gestao_projetos\execucao_financeira"
fin_file = os.path.join(base_dir, "execucao_financeira.xlsx")
conv_dir = os.path.join(base_dir, "dados-conveniar")

wb_fin = openpyxl.load_workbook(fin_file, data_only=True)
ws_ext = wb_fin["extrato_bancario"]

conta_files = {
    '375 - IFAM 2026-006 - Bio Caroço (15334-6)': '1793138381_ExtratoProjeto_375.xlsx',
    '376 - IFAM 2026-006 - Bio Caroço SEBRAE (15335-4)': '1793138497_ExtratoProjeto_376.xlsx',
    '377 - IFAM 2026-006 - Bio Caroço EMBRAPII (15338-9)': '1793138518_ExtratoProjeto_377.xlsx',
}

for conta, fname in conta_files.items():
    print("=" * 80)
    print(f"CONTA: {conta}")
    print(f"ARQUIVO CONVENIAR: {fname}")
    print("=" * 80)
    
    # 1. Registros na planilha atual para essa conta
    existentes = []
    for r in range(2, ws_ext.max_row + 1):
        c_val = str(ws_ext.cell(r, 4).value or "").strip()
        if c_val and len(c_val.split()) > 0 and c_val.split()[0] == conta.split()[0]:
            existentes.append({
                'row': r,
                'lan': ws_ext.cell(r, 1).value,
                'pag': ws_ext.cell(r, 2).value,
                'mes': ws_ext.cell(r, 3).value,
                'rub': ws_ext.cell(r, 5).value,
                'doc': str(ws_ext.cell(r, 6).value or "").strip(),
                'mov': ws_ext.cell(r, 7).value,
                'tipo': ws_ext.cell(r, 8).value,
                'desc': ws_ext.cell(r, 9).value,
                'val': ws_ext.cell(r, 10).value
            })
            
    # 2. Lançamentos no arquivo Conveniar
    wb_c = openpyxl.load_workbook(os.path.join(conv_dir, fname), data_only=True)
    ws_c = wb_c[wb_c.sheetnames[0]]
    conv_items = []
    current_rub = ""
    for r in range(1, ws_c.max_row + 1):
        col_b = str(ws_c.cell(r, 2).value or "").strip()
        if col_b.startswith("Rubrica:"):
            current_rub = col_b.replace("Rubrica:", "").strip()
            continue
        dt_lan = ws_c.cell(r, 2).value
        if not dt_lan or str(dt_lan).startswith("Rubrica") or str(dt_lan).startswith("Data") or str(dt_lan).startswith("Saldo") or str(dt_lan).startswith("Gerado") or str(dt_lan).startswith("Devolu") or str(dt_lan).startswith("Total"):
            continue
        hist = str(ws_c.cell(r, 6).value or "").strip()
        if hist.startswith("Total da Rubrica") or hist.startswith("Saldo Final"):
            continue
            
        deb = ws_c.cell(r, 7).value
        cred = ws_c.cell(r, 8).value
        if deb is not None and str(deb).strip() != "":
            val = float(deb)
            if val > 0: val = -val
            mov = "debito"
        elif cred is not None and str(cred).strip() != "":
            val = abs(float(cred))
            mov = "credito"
        else:
            continue
            
        conv_items.append({
            'source_row': r,
            'lan': dt_lan,
            'pag': ws_c.cell(r, 3).value,
            'doc': str(ws_c.cell(r, 4).value or "").strip(),
            'tipo': str(ws_c.cell(r, 5).value or "").strip(),
            'desc': hist,
            'rub': current_rub,
            'val': val,
            'mov': mov
        })

    print(f"Existentes na planilha: {len(existentes)} registros")
    print(f"Lançamentos no Conveniar: {len(conv_items)} registros")
    
    # Imprimir todos os lançamentos do Conveniar e se já existem ou não
    print("\n--- Lançamentos do Conveniar para esta conta ---")
    for ci in conv_items:
        match = [e for e in existentes if e['doc'] == ci['doc'] and abs(float(e['val']) - float(ci['val'])) < 0.01]
        status = f"JA EXISTE (Linha {match[0]['row']})" if match else "** NOVO **"
        print(f"  [{status}] Doc: {ci['doc']:12} | Pag: {ci['pag']} | Val: {ci['val']:10.2f} | Rub: {ci['rub']} | Tipo: {ci['tipo']}")
        if not match:
            print(f"      -> Desc: {ci['desc']}")
            
    # Verificar se há registros existentes de setembro que talvez foram ajustados
    docs_conv = set(ci['doc'] for ci in conv_items)
    print("\n--- Registros existentes na planilha em SETEMBRO/OUTUBRO que NÃO estão no arquivo Conveniar ---")
    for e in existentes:
        pag_str = str(e['pag'])
        if ('2026-09' in pag_str or 'setembro' in str(e['mes']) or '2026-10' in pag_str or 'outubro' in str(e['mes'])) and e['doc'] not in docs_conv:
            print(f"  L{e['row']}: Doc={e['doc']} | Pag={e['pag']} | Val={e['val']} | Rub={e['rub']} | Tipo={e['tipo']} | Desc={e['desc'][:70]}")
