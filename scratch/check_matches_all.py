import os
import openpyxl

base_dir = r"C:\Projetos\ARGUS\gestao_projetos\execucao_financeira"
fin_file = os.path.join(base_dir, "execucao_financeira.xlsx")
conv_dir = os.path.join(base_dir, "dados-conveniar")

wb_fin = openpyxl.load_workbook(fin_file, data_only=True)
ws_ext = wb_fin["extrato_bancario"]

conta_files = {
    '375 - IFAM 2026-006 - Bio Caroço (15334-6)': '1793138381_ExtratoProjeto_375.xlsx',
    '376 - IFAM 2026-006 - Bio Caroço SEBRAE (15335-4)': '1793138497_ExtratoProjeto_376.xlsx',
    '377 - IFAM 2026-006 - Bio Caroço EMBRAPII (15338-9)': '1793138518_ExtratoProjeto_377.xlsx',
}

for conta, fname in conta_files.items():
    print("=" * 80)
    print(f"CONTA: {conta} ({fname})")
    print("=" * 80)
    
    # ler conveniar
    wb_c = openpyxl.load_workbook(os.path.join(conv_dir, fname), data_only=True)
    ws_c = wb_c[wb_c.sheetnames[0]]
    current_rub = ""
    
    for r in range(1, ws_c.max_row + 1):
        col_b = str(ws_c.cell(r, 2).value or "").strip()
        if col_b.startswith("Rubrica:"):
            current_rub = col_b.replace("Rubrica:", "").strip()
            continue
        dt_lan = ws_c.cell(r, 2).value
        if not dt_lan or str(dt_lan).startswith("Rubrica") or str(dt_lan).startswith("Data") or str(dt_lan).startswith("Saldo") or str(dt_lan).startswith("Gerado") or str(dt_lan).startswith("Devolu") or str(dt_lan).startswith("Total"):
            continue
        hist = str(ws_c.cell(r, 6).value or "").strip()
        if hist.startswith("Total da Rubrica") or hist.startswith("Saldo Final"):
            continue
            
        deb = ws_c.cell(r, 7).value
        cred = ws_c.cell(r, 8).value
        if deb is not None and str(deb).strip() != "":
            val = float(deb)
            if val > 0: val = -val
            mov = "debito"
        elif cred is not None and str(cred).strip() != "":
            val = abs(float(cred))
            mov = "credito"
        else:
            continue
            
        dt_pag = ws_c.cell(r, 3).value
        doc = str(ws_c.cell(r, 4).value or "").strip()
        tipo = str(ws_c.cell(r, 5).value or "").strip()
        
        # Procurar correspondente na planilha atual
        matches = []
        for row_idx in range(2, ws_ext.max_row + 1):
            e_c = str(ws_ext.cell(row_idx, 4).value or "").strip()
            if e_c and e_c.split()[0] == conta.split()[0]:
                e_doc = str(ws_ext.cell(row_idx, 6).value or "").strip()
                e_val = ws_ext.cell(row_idx, 10).value
                if e_val is not None:
                    if e_doc == doc and abs(float(e_val) - val) < 0.01:
                        matches.append(row_idx)
        
        status = f"JA EXISTE (Linha {matches[0]})" if matches else "** NOVO LANÇAMENTO **"
        print(f"R{r:2d}: {status} | Pag:{dt_pag} | Doc:{doc:12} | Val:{val:10.2f} | Rub:{current_rub:25} | Tipo:{tipo}")
