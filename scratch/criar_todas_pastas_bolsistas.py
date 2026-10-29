import openpyxl
import os

excel_path = 'gestao_projetos/media/modelos/relatorios-atividades-cabecalhos.xlsx'
wb = openpyxl.load_workbook(excel_path)
ws = wb.active
headers = [ws.cell(1, c).value for c in range(1, ws.max_column+1)]
idx_nome = headers.index('bolsista_nome')

bolsistas = set()
for r in range(2, ws.max_row+1):
    val = ws.cell(r, idx_nome+1).value
    if val and str(val).strip():
        bolsistas.add(str(val).strip())

base_dir = os.path.abspath('gestao_projetos/media/modelos')
chars_proibidos = ['\\', '/', '*', '?', ':', '"', '<', '>', '|']

criadas = 0
for b in sorted(bolsistas):
    pasta_nome = b
    for ch in chars_proibidos:
        pasta_nome = pasta_nome.replace(ch, '_')
    dir_path = os.path.join(base_dir, pasta_nome)
    if not os.path.exists(dir_path):
        os.makedirs(dir_path, exist_ok=True)
        print(f"+ Criada pasta: {pasta_nome}")
        criadas += 1
    else:
        print(f"= Ja existia: {pasta_nome}")

print(f"\nTotal de bolsistas no projeto: {len(bolsistas)}")
print(f"Novas pastas criadas: {criadas}")
