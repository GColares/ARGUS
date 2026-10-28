import os
import openpyxl

base_dir = r"C:\Projetos\ARGUS\gestao_projetos\execucao_financeira"
fin_file = os.path.join(base_dir, "execucao_financeira.xlsx")
conv_dir = os.path.join(base_dir, "dados-conveniar")

wb_fin = openpyxl.load_workbook(fin_file, data_only=True)
ws_ext = wb_fin["extrato_bancario"]

# Inspecionar registros existentes para ver formato de mes, tipos, rubricas
print("Amostra de 10 registros de extrato_bancario:")
for r in range(2, 12):
    row_vals = [ws_ext.cell(r, c).value for c in range(1, 11)]
    print(f"L{r}: {row_vals}")

# Verificar formato da coluna 'mês'
meses = set(ws_ext.cell(r, 3).value for r in range(2, ws_ext.max_row + 1) if ws_ext.cell(r, 3).value)
print("Valores únicos na coluna 'mês':", meses)

# Verificar se 'mês' é baseado em data_pagamento ou data_lancamento
divergencias_mes = []
mes_map = {1: 'janeiro', 2: 'fevereiro', 3: 'março', 4: 'abril', 5: 'maio', 6: 'junho',
           7: 'julho', 8: 'agosto', 9: 'setembro', 10: 'outubro', 11: 'novembro', 12: 'dezembro'}

for r in range(2, ws_ext.max_row + 1):
    dt_lan = ws_ext.cell(r, 1).value
    dt_pag = ws_ext.cell(r, 2).value
    m_val = ws_ext.cell(r, 3).value
    if not m_val: continue
    
    # testar mes de dt_lan vs dt_pag
    m_lan = mes_map.get(dt_lan.month) if hasattr(dt_lan, 'month') else None
    m_pag = mes_map.get(dt_pag.month) if hasattr(dt_pag, 'month') else None
    if m_lan != m_pag:
        divergencias_mes.append((r, dt_lan, dt_pag, m_val, f"lan={m_lan}, pag={m_pag}"))

print(f"Total divergências entre mês de lançamento e pagamento: {len(divergencias_mes)}")
for d in divergencias_mes[:5]:
    print("  ", d)
