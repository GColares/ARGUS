import os
import shutil
import openpyxl

# Caminhos base
origem_modelos = os.path.abspath('gestao_projetos/media/modelos')
destino_projeto = os.path.abspath('gestao_projetos/projetos/bio_caroco')
destino_bolsistas = os.path.join(destino_projeto, 'bolsistas')

os.makedirs(destino_bolsistas, exist_ok=True)
print(f"Diretório base do projeto configurado: {destino_projeto}")
print(f"Diretório de bolsistas configurado: {destino_bolsistas}")

# 1. Carregar todos os bolsistas da planilha de cabeçalhos
excel_path = os.path.join(origem_modelos, 'relatorios-atividades-cabecalhos.xlsx')
wb = openpyxl.load_workbook(excel_path)
ws = wb.active
headers = [ws.cell(1, c).value for c in range(1, ws.max_column+1)]
idx_nome = headers.index('bolsista_nome')

bolsistas = set()
for r in range(2, ws.max_row+1):
    val = ws.cell(r, idx_nome+1).value
    if val and str(val).strip():
        bolsistas.add(str(val).strip())

chars_proibidos = ['\\', '/', '*', '?', ':', '"', '<', '>', '|']

# 2. Criar ou mover as pastas de todos os 12 bolsistas para o novo destino
pastas_migradas = 0
pastas_criadas = 0

for b in sorted(bolsistas):
    pasta_nome = b
    for ch in chars_proibidos:
        pasta_nome = pasta_nome.replace(ch, '_')
        
    pasta_antiga = os.path.join(origem_modelos, pasta_nome)
    pasta_nova = os.path.join(destino_bolsistas, pasta_nome)
    
    if os.path.exists(pasta_antiga) and os.path.isdir(pasta_antiga):
        # Mover a pasta com todos os seus arquivos dentro
        if os.path.exists(pasta_nova):
            # Se já existir, mover arquivos internos
            for item in os.listdir(pasta_antiga):
                s = os.path.join(pasta_antiga, item)
                d = os.path.join(pasta_nova, item)
                shutil.move(s, d)
            os.rmdir(pasta_antiga)
        else:
            shutil.move(pasta_antiga, pasta_nova)
        print(f"Migrada pasta e relatorios: {pasta_nome} -> {destino_bolsistas}")
        pastas_migradas += 1
    else:
        os.makedirs(pasta_nova, exist_ok=True)
        print(f"Criada pasta nova: {pasta_nome} em {destino_bolsistas}")
        pastas_criadas += 1

print(f"\nConcluido! {pastas_migradas} pastas migradas com arquivos e {pastas_criadas} novas pastas criadas.")
