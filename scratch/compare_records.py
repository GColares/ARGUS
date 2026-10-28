import os
import openpyxl

base_dir = r"C:\Projetos\ARGUS\gestao_projetos\execucao_financeira"
fin_file = os.path.join(base_dir, "execucao_financeira.xlsx")
conv_dir = os.path.join(base_dir, "dados-conveniar")

wb_fin = openpyxl.load_workbook(fin_file, data_only=True)
ws_ext = wb_fin["extrato_bancario"]

existing_records = []
for r in range(2, ws_ext.max_row + 1):
    dt_lan = ws_ext.cell(r, 1).value
    dt_pag = ws_ext.cell(r, 2).value
    mes = ws_ext.cell(r, 3).value
    conta = ws_ext.cell(r, 4).value
    rub = ws_ext.cell(r, 5).value
    doc = ws_ext.cell(r, 6).value
    mov = ws_ext.cell(r, 7).value
    tipo = ws_ext.cell(r, 8).value
    desc = ws_ext.cell(r, 9).value
    val = ws_ext.cell(r, 10).value
    
    if not any([dt_lan, dt_pag, conta, val]): continue
    
    existing_records.append({
        'row': r, 'dt_lan': dt_lan, 'dt_pag': dt_pag, 'mes': mes,
        'conta': str(conta or "").strip(),
        'rubrica': str(rub or "").strip(),
        'doc': str(doc or "").strip(),
        'mov': str(mov or "").strip(),
        'tipo': str(tipo or "").strip(),
        'desc': str(desc or "").strip(),
        'val': round(float(val), 2) if val is not None else 0.0
    })

print(f"Total registros existentes em execucao_financeira.xlsx: {len(existing_records)}")

# Mapeamento do nome do arquivo para o nome canônico da Conta
conta_map = {
    '1793138381_ExtratoProjeto_375.xlsx': '375 - IFAM 2026-006 - Bio Caroço (15334-6)',
    '1793138497_ExtratoProjeto_376.xlsx': '376 - IFAM 2026-006 - Bio Caroço SEBRAE (15335-4)',
    '1793138518_ExtratoProjeto_377.xlsx': '377 - IFAM 2026-006 - Bio Caroço EMBRAPII (15338-9)',
}

conveniar_records = []

for fname, conta_nome in conta_map.items():
    fpath = os.path.join(conv_dir, fname)
    wb_c = openpyxl.load_workbook(fpath, data_only=True)
    ws_c = wb_c[wb_c.sheetnames[0]]
    current_rubrica = ""
    
    for r in range(1, ws_c.max_row + 1):
        col_b = str(ws_c.cell(r, 2).value or "").strip()
        if col_b.startswith("Rubrica:"):
            current_rubrica = col_b.replace("Rubrica:", "").strip()
            continue
            
        dt_lan = ws_c.cell(r, 2).value
        # checar se é data válida
        if not dt_lan or str(dt_lan).startswith("Rubrica") or str(dt_lan).startswith("Data") or str(dt_lan).startswith("Saldo") or str(dt_lan).startswith("Gerado"):
            continue
            
        dt_pag = ws_c.cell(r, 3).value
        doc = ws_c.cell(r, 4).value
        tipo = ws_c.cell(r, 5).value
        hist = ws_c.cell(r, 6).value
        deb = ws_c.cell(r, 7).value
        cred = ws_c.cell(r, 8).value
        
        # Ignorar linhas de totais
        if str(hist).startswith("Total da Rubrica") or str(hist).startswith("Saldo Final"):
            continue
            
        # Determinar valor e movimento
        if deb is not None and str(deb).strip() != "":
            # debito no conveniar pode vir negativo ou positivo
            val_num = float(deb)
            if val_num > 0: val_num = -val_num
            mov = "debito"
        elif cred is not None and str(cred).strip() != "":
            val_num = abs(float(cred))
            mov = "credito"
        else:
            continue
            
        conveniar_records.append({
            'file': fname,
            'source_row': r,
            'conta': conta_nome,
            'rubrica': current_rubrica,
            'dt_lan': dt_lan,
            'dt_pag': dt_pag,
            'doc': str(doc or "").strip(),
            'tipo': str(tipo or "").strip(),
            'desc': str(hist or "").strip(),
            'mov': mov,
            'val': round(val_num, 2)
        })

print(f"Total registros válidos lidos do Conveniar: {len(conveniar_records)}")

# Vamos checar a correspondência
matched = 0
unmatched_new = []
for cr in conveniar_records:
    # procurar match
    found = False
    for er in existing_records:
        # Match por doc, conta e valor (ou doc e conta)
        if cr['doc'] and er['doc'] == cr['doc'] and er['conta'].split()[0] == cr['conta'].split()[0] and abs(er['val'] - cr['val']) < 0.01:
            found = True
            break
        elif cr['doc'] and er['doc'] == cr['doc'] and er['conta'].split()[0] == cr['conta'].split()[0]:
            print(f"[DOC MATCH MAS VALOR DIFERENTE] Conv Doc: {cr['doc']}, Conv Val: {cr['val']} vs Exist Val: {er['val']}")
            found = True
            break
    if found:
        matched += 1
    else:
        unmatched_new.append(cr)

print(f"\nRegistros do Conveniar que já existem na planilha: {matched}")
print(f"Registros do Conveniar que são NOVOS: {len(unmatched_new)}")

print("\n--- DETALHES DOS REGISTROS NOVOS ---")
for i, nr in enumerate(unmatched_new, 1):
    print(f"{i}. Conta: {nr['conta'][:25]} | Rub: {nr['rubrica']} | Pag: {nr['dt_pag']} | Doc: {nr['doc']} | Val: {nr['val']} | Tipo: {nr['tipo']}")
    print(f"   Desc: {nr['desc'][:90]}")
