import os
import django
import io

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'argus_core.settings')
django.setup()

from django.test import Client
from cadastros.models import ProjetoPDI, AtividadePlanoAcao

# Find a project
p = ProjetoPDI.objects.first()

client = Client()

# Create a fake docx file
from docx import Document
doc = Document()
doc.add_paragraph("Cronograma de Atividades")
doc.add_paragraph("Item 1")
doc.add_paragraph("Nome da Atividade: Teste Atividade")
doc.add_paragraph("Descrição da Atividade: Teste Desc")
doc.add_paragraph("Justificativa: Teste Just")
doc.add_paragraph("Entregáveis: Teste Ent")
doc.add_paragraph("Data de início: Mês 1")
doc.add_paragraph("Data de fim: Mês 2")

f = io.BytesIO()
doc.save(f)
f.seek(0)
f.name = 'teste.docx'

print("Uploading file...")
response = client.post(f'/projeto/{p.id}/importar-cronograma/', {'arquivo_docx': f})
print("Upload response:", response.status_code)

if response.status_code == 200:
    print("Parsing form data...")
    # Get form data from response context
    formset = response.context['formset']
    post_data = {
        'form-TOTAL_FORMS': formset.management_form['TOTAL_FORMS'].value(),
        'form-INITIAL_FORMS': formset.management_form['INITIAL_FORMS'].value(),
        'form-MIN_NUM_FORMS': formset.management_form['MIN_NUM_FORMS'].value(),
        'form-MAX_NUM_FORMS': formset.management_form['MAX_NUM_FORMS'].value(),
    }
    
    for i, form in enumerate(formset.forms):
        for name, field in form.fields.items():
            val = form[name].value()
            post_data[f'form-{i}-{name}'] = val if val is not None else ''

    print("Submitting form data...")
    response2 = client.post(f'/projeto/{p.id}/importar-cronograma/', post_data)
    print("Submit response:", response2.status_code)
    
    if response2.status_code == 302:
        print("Count after submit:", AtividadePlanoAcao.objects.filter(projeto=p).count())
        print("Items:", list(AtividadePlanoAcao.objects.filter(projeto=p).values('nome')))
    else:
        print("Failed to submit. Form errors:")
        print(response2.context['formset'].errors)
