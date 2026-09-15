# -*- coding: utf-8 -*-
"""
Exportador de Planilha Excel Consolidada (OpenPyXL).
Gera as 5 abas relacionais com base nos dados reais do SQLite:
1. Catálogo Consolidado (com Saldo)
2. Histórico de Entradas (Compras/NFs)
3. Histórico de Saídas (Consumo)
4. Empréstimos de Ferramentas
5. Parâmetros do Sistema
"""
import os
import shutil
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from banco import get_connection, get_catalogo, DB_PATH

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DIR_EXCEL = os.path.join(BASE_DIR, "planilhas_consolidadas")
DIR_DRIVE = os.path.join(BASE_DIR, "backups_drive")

os.makedirs(DIR_EXCEL, exist_ok=True)
os.makedirs(DIR_DRIVE, exist_ok=True)

VERDE_TOPO = "004D25"
VERDE_HEADER = "006633"
AZUL_HEADER = "0D47A1"
BRANCO = "FFFFFF"
ZEBRA = "F9FAFB"
BORDA = "D1D5DB"

def exportar_planilha_completa(caminho_custom_drive=None):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    font_tit = Font(name="Calibri", size=13, bold=True, color=BRANCO)
    font_head = Font(name="Calibri", size=10, bold=True, color=BRANCO)
    font_norm = Font(name="Calibri", size=10)
    fill_topo = PatternFill(start_color=VERDE_TOPO, end_color=VERDE_TOPO, fill_type="solid")
    fill_verde = PatternFill(start_color=VERDE_HEADER, end_color=VERDE_HEADER, fill_type="solid")
    fill_azul = PatternFill(start_color=AZUL_HEADER, end_color=AZUL_HEADER, fill_type="solid")
    fill_zebra = PatternFill(start_color=ZEBRA, end_color=ZEBRA, fill_type="solid")

    borda_fina = Border(
        left=Side(style='thin', color=BORDA), right=Side(style='thin', color=BORDA),
        top=Side(style='thin', color=BORDA), bottom=Side(style='thin', color=BORDA)
    )

    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')

    conn = get_connection(DB_PATH)

    # 1. ABA: TB_CATALOGO_MATERIAIS
    ws1 = wb.create_sheet(title="TB_CATALOGO_MATERIAIS")
    ws1.views.sheetView[0].showGridLines = True
    ws1.merge_cells("A1:K1")
    ws1["A1"] = "CATÁLOGO DE MATERIAIS & FERRAMENTAS — ESTOQUE ATUALIZADO"
    ws1["A1"].font = font_tit
    ws1["A1"].fill = fill_topo
    ws1["A1"].alignment = align_center
    ws1.row_dimensions[1].height = 28

    headers1 = [
        "Código", "Descrição do Material / Ferramenta", "Categoria",
        "Unid.", "Estoque Inicial", "Total Entradas", "Total Saídas",
        "Saldo Atual", "Est. Mínimo", "Status", "Localização"
    ]
    ws1.row_dimensions[3].height = 22
    for c_idx, h in enumerate(headers1, 1):
        cell = ws1.cell(row=3, column=c_idx, value=h)
        cell.font = font_head
        cell.fill = fill_verde
        cell.alignment = align_center
        cell.border = borda_fina

    materiais = get_catalogo(db_path=DB_PATH, apenas_ativos=False)
    for r_idx, m in enumerate(materiais, 4):
        ws1.row_dimensions[r_idx].height = 19
        ws1.cell(row=r_idx, column=1, value=m['id']).alignment = align_center
        ws1.cell(row=r_idx, column=2, value=m['descricao']).alignment = align_left
        ws1.cell(row=r_idx, column=3, value=m['categoria']).alignment = align_center
        ws1.cell(row=r_idx, column=4, value=m['unidade']).alignment = align_center
        ws1.cell(row=r_idx, column=5, value=m['estoque_inicial']).alignment = align_right
        ws1.cell(row=r_idx, column=6, value=m['total_entradas']).alignment = align_right
        ws1.cell(row=r_idx, column=7, value=m['total_saidas']).alignment = align_right
        ws1.cell(row=r_idx, column=8, value=m['saldo']).alignment = align_right
        ws1.cell(row=r_idx, column=9, value=m['estoque_minimo']).alignment = align_right
        ws1.cell(row=r_idx, column=10, value="🔴 BAIXO" if m['status_estoque'] == "BAIXO" else "🟢 REGULAR").alignment = align_center
        ws1.cell(row=r_idx, column=11, value=m['localizacao']).alignment = align_left

        for col in range(1, 12):
            c = ws1.cell(row=r_idx, column=col)
            c.font = font_norm
            c.border = borda_fina
            if r_idx % 2 == 0:
                c.fill = fill_zebra

    for col in ws1.columns:
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = 16
    ws1.column_dimensions['B'].width = 45

    # 2. ABA: TB_HISTORICO_ENTRADAS
    ws2 = wb.create_sheet(title="TB_HISTORICO_ENTRADAS")
    ws2.views.sheetView[0].showGridLines = True
    ws2.merge_cells("A1:K1")
    ws2["A1"] = "HISTÓRICO DE ENTRADAS DE MATERIAIS (COMPRAS / NFS / DEVOLUÇÕES)"
    ws2["A1"].font = font_tit
    ws2["A1"].fill = fill_topo
    ws2["A1"].alignment = align_center
    ws2.row_dimensions[1].height = 28

    headers2 = ["ID", "Data/Hora", "Nº NF", "Fornecedor / Origem", "Material", "Qtd Compra", "Unid Compra", "Qtd Base", "Unid Base", "Valor Unit (R$)", "Valor Total (R$)", "Projeto"]
    ws2.row_dimensions[3].height = 22
    for c_idx, h in enumerate(headers2, 1):
        cell = ws2.cell(row=3, column=c_idx, value=h)
        cell.font = font_head
        cell.fill = fill_verde
        cell.alignment = align_center
        cell.border = borda_fina

    cur = conn.cursor()
    cur.execute("""
    SELECT e.id, e.data_hora, e.nf, e.fornecedor, m.descricao, e.quantidade_compra, e.unidade_compra, e.quantidade_base, m.unidade_base, e.valor_unit, e.valor_total, e.projeto
    FROM tb_entradas e
    JOIN tb_materiais m ON e.material_id = m.id
    ORDER BY e.id DESC
    """)
    entradas = cur.fetchall()
    for r_idx, ent in enumerate(entradas, 4):
        ws2.row_dimensions[r_idx].height = 19
        for c_idx, val in enumerate(ent, 1):
            c = ws2.cell(row=r_idx, column=c_idx, value=val)
            c.font = font_norm
            c.border = borda_fina
            c.alignment = align_center if c_idx in [1, 2, 3, 6, 7, 8, 9] else (align_right if c_idx in [10, 11] else align_left)
            if r_idx % 2 == 0:
                c.fill = fill_zebra

    for col in ws2.columns:
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = 16
    ws2.column_dimensions['E'].width = 38

    # 3. ABA: TB_HISTORICO_SAIDAS
    ws3 = wb.create_sheet(title="TB_HISTORICO_SAIDAS")
    ws3.views.sheetView[0].showGridLines = True
    ws3.merge_cells("A1:J1")
    ws3["A1"] = "HISTÓRICO DE SAÍDAS DE CONSUMO E APLICAÇÃO"
    ws3["A1"].font = font_tit
    ws3["A1"].fill = fill_topo
    ws3["A1"].alignment = align_center
    ws3.row_dimensions[1].height = 28

    headers3 = ["ID Saída", "Data/Hora", "Item", "Qtd Base", "Unid Base", "Solicitante", "Projeto", "Ambiente", "Tipo", "Recibo"]
    ws3.row_dimensions[3].height = 22
    for c_idx, h in enumerate(headers3, 1):
        cell = ws3.cell(row=3, column=c_idx, value=h)
        cell.font = font_head
        cell.fill = fill_verde
        cell.alignment = align_center
        cell.border = borda_fina

    cur.execute("SELECT id, data_hora, descricao, quantidade_base, unidade_base, solicitante, projeto, ambiente, tipo_saida, comprovante_num FROM tb_saidas ORDER BY id DESC")
    saidas = cur.fetchall()
    for r_idx, s in enumerate(saidas, 4):
        ws3.row_dimensions[r_idx].height = 19
        for c_idx, val in enumerate(s, 1):
            c = ws3.cell(row=r_idx, column=c_idx, value=val)
            c.font = font_norm
            c.border = borda_fina
            c.alignment = align_center if c_idx in [1, 2, 4, 5, 9, 10] else align_left
            if r_idx % 2 == 0:
                c.fill = fill_zebra

    for col in ws3.columns:
        col_letter = get_column_letter(col[0].column)
        ws3.column_dimensions[col_letter].width = 18
    ws3.column_dimensions['C'].width = 38

    # 4. ABA: TB_EMPRESTIMOS_FERRAMENTAS
    ws4 = wb.create_sheet(title="TB_EMPRESTIMOS_FERRAMENTAS")
    ws4.views.sheetView[0].showGridLines = True
    ws4.merge_cells("A1:I1")
    ws4["A1"] = "CONTROLE DE EMPRÉSTIMOS E DEVOLUÇÃO DE FERRAMENTAS"
    ws4["A1"].font = font_tit
    ws4["A1"].fill = fill_topo
    ws4["A1"].alignment = align_center
    ws4.row_dimensions[1].height = 28

    headers4 = ["ID Empréstimo", "Ferramenta", "Solicitante", "Projeto", "Data Saída", "Prev. Devolução", "Status", "Data Retorno", "Avarias/Obs"]
    ws4.row_dimensions[3].height = 22
    for c_idx, h in enumerate(headers4, 1):
        cell = ws4.cell(row=3, column=c_idx, value=h)
        cell.font = font_head
        cell.fill = fill_azul
        cell.alignment = align_center
        cell.border = borda_fina

    cur.execute("SELECT id, ferramenta_nome, solicitante, projeto, data_saida, data_prevista, status, data_retorno, avarias FROM tb_emprestimos_ferramentas ORDER BY id DESC")
    emps = cur.fetchall()
    for r_idx, emp in enumerate(emps, 4):
        ws4.row_dimensions[r_idx].height = 19
        for c_idx, val in enumerate(emp, 1):
            c = ws4.cell(row=r_idx, column=c_idx, value=val)
            c.font = font_norm
            c.border = borda_fina
            c.alignment = align_center if c_idx in [1, 5, 6, 7, 8] else align_left
            if r_idx % 2 == 0:
                c.fill = fill_zebra

    for col in ws4.columns:
        col_letter = get_column_letter(col[0].column)
        ws4.column_dimensions[col_letter].width = 20
    ws4.column_dimensions['B'].width = 35

    # 5. ABA: TB_PARAMETROS
    ws5 = wb.create_sheet(title="TB_PARAMETROS")
    ws5.views.sheetView[0].showGridLines = True
    ws5.merge_cells("A1:C1")
    ws5["A1"] = "PARÂMETROS CADASTRADOS NO SISTEMA (ARGUS)"
    ws5["A1"].font = font_tit
    ws5["A1"].fill = fill_topo
    ws5["A1"].alignment = align_center
    ws5.row_dimensions[1].height = 28

    ws5["A3"] = "SOLICITANTE (NOME)"
    ws5["B3"] = "SOLICITANTE (CPF)"
    ws5["C3"] = "PROJETOS"
    ws5["D3"] = "AMBIENTES"
    for c_idx, col in enumerate(["A", "B", "C", "D"], 1):
        ws5[f"{col}3"].font = font_head
        ws5[f"{col}3"].fill = fill_verde
        ws5[f"{col}3"].alignment = align_center

    cur.execute("SELECT nome_completo, COALESCE(cpf, '') FROM tb_solicitantes WHERE ativo = 1 ORDER BY nome_completo ASC")
    sols = cur.fetchall()
    cur.execute("SELECT valor FROM tb_parametros WHERE tipo = 'PROJETO' ORDER BY valor ASC")
    projs = [r[0] for r in cur.fetchall()]
    cur.execute("SELECT valor FROM tb_parametros WHERE tipo = 'AMBIENTE' ORDER BY valor ASC")
    ambs = [r[0] for r in cur.fetchall()]

    max_len = max(len(sols), len(projs), len(ambs), 1)
    for r in range(max_len):
        row_num = r + 4
        if r < len(sols):
            ws5.cell(row=row_num, column=1, value=sols[r][0]).font = font_norm
            ws5.cell(row=row_num, column=2, value=sols[r][1]).font = font_norm
        if r < len(projs): ws5.cell(row=row_num, column=3, value=projs[r]).font = font_norm
        if r < len(ambs): ws5.cell(row=row_num, column=4, value=ambs[r]).font = font_norm

    ws5.column_dimensions['A'].width = 38
    ws5.column_dimensions['B'].width = 20
    ws5.column_dimensions['C'].width = 35
    ws5.column_dimensions['D'].width = 35

    conn.close()

    caminho_consolidado = os.path.join(DIR_EXCEL, "Almoxarifado_ARGUS_Consolidado.xlsx")
    wb.save(caminho_consolidado)

    data_str = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    nome_backup = f"Backup_Almoxarifado_{data_str}.xlsx"
    caminho_backup_drive = os.path.join(DIR_DRIVE, nome_backup)
    wb.save(caminho_backup_drive)

    # Cópia do banco de dados SQLite local
    nome_backup_db = f"Backup_Banco_{data_str}.db"
    caminho_backup_db_local = os.path.join(DIR_DRIVE, nome_backup_db)
    if os.path.exists(DB_PATH):
        try:
            shutil.copy2(DB_PATH, caminho_backup_db_local)
        except Exception as e:
            print(f"Aviso ao copiar DB local: {e}")

    copia_custom = None
    if caminho_custom_drive and os.path.exists(caminho_custom_drive):
        try:
            # 1. Copia a planilha consolidada oficial atualizada
            shutil.copy2(caminho_consolidado, os.path.join(caminho_custom_drive, "Almoxarifado_ARGUS_Consolidado.xlsx"))
            # 2. Copia o backup histórico do Excel
            copia_custom = os.path.join(caminho_custom_drive, nome_backup)
            wb.save(copia_custom)
            # 3. Copia a segurança do banco SQLite
            if os.path.exists(DB_PATH):
                shutil.copy2(DB_PATH, os.path.join(caminho_custom_drive, nome_backup_db))
        except Exception as e:
            print(f"Erro ao copiar para caminho customizado do Drive: {e}")

    return {
        "arquivo_consolidado": caminho_consolidado,
        "arquivo_backup": caminho_backup_drive,
        "arquivo_drive_custom": copia_custom
    }

