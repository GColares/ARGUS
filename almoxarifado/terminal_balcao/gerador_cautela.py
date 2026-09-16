# -*- coding: utf-8 -*-
"""
Gerador de Termo de Cautela e Responsabilidade de Equipamento (ARGUS Almoxarifado).
Carrega o template DOCX oficial ('proposta-termo-cautela-equipamento.docx') e
injeta com precisao os dados do documento, partes envolvidas e lista dinamica
de equipamentos na tabela.
"""
import os
import sys
import re
from datetime import datetime
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
MODELOS_DIR = os.path.abspath(os.path.join(BASE_DIR, "..", "..", "cadastros", "modelos"))
TEMPLATE_DOCX = os.path.join(MODELOS_DIR, "proposta-termo-cautela-equipamento.docx")

if not os.path.exists(TEMPLATE_DOCX):
    alt = os.path.join(BASE_DIR, "proposta-termo-cautela-equipamento.docx")
    if os.path.exists(alt):
        TEMPLATE_DOCX = alt

DIR_SAIDA_CAUTELAS = os.path.join(BASE_DIR, "cautelas_geradas")
os.makedirs(DIR_SAIDA_CAUTELAS, exist_ok=True)

MESES_PTBR = {
    1: "janeiro", 2: "fevereiro", 3: "março", 4: "abril",
    5: "maio", 6: "junho", 7: "julho", 8: "agosto",
    9: "setembro", 10: "outubro", 11: "novembro", 12: "dezembro"
}

def formatar_data_extenso(data_str=None):
    if not data_str:
        dt = datetime.now()
    else:
        d_str = str(data_str).strip()
        try:
            if "-" in d_str:
                dt = datetime.strptime(d_str[:10], "%Y-%m-%d")
            else:
                dt = datetime.strptime(d_str[:10], "%d/%m/%Y")
        except Exception:
            dt = datetime.now()

    dia = dt.day
    mes = MESES_PTBR.get(dt.month, "")
    ano = dt.year
    return f"{dia} de {mes} de {ano}"

def formatar_data_br(data_str=None):
    if not data_str:
        return datetime.now().strftime("%d/%m/%Y")
    d_str = str(data_str).strip()
    try:
        if "-" in d_str:
            dt = datetime.strptime(d_str[:10], "%Y-%m-%d")
            return dt.strftime("%d/%m/%Y")
        return d_str
    except Exception:
        return d_str

def sanitizar_nome_arquivo(nome):
    s = re.sub(r'[^a-zA-Z0-9_\-\.]', '_', str(nome).strip())
    return s[:60]

def substituir_texto_preservando_estilo(p, novo_texto):
    fonte_nome = None
    fonte_tam = None
    fonte_negrito = None
    fonte_cor = None

    if p.runs:
        r0 = p.runs[0]
        fonte_nome = r0.font.name
        fonte_tam = r0.font.size
        fonte_negrito = r0.bold
        if r0.font.color and r0.font.color.rgb:
            fonte_cor = r0.font.color.rgb

    p.text = novo_texto

    if p.runs and (fonte_nome or fonte_tam or fonte_negrito is not None or fonte_cor):
        for r in p.runs:
            if fonte_nome:
                r.font.name = fonte_nome
            if fonte_tam:
                r.font.size = fonte_tam
            if fonte_negrito is not None:
                r.bold = fonte_negrito
            if fonte_cor:
                r.font.color.rgb = fonte_cor

def gerar_termo_cautela_docx(dados, caminho_saida=None):
    if not os.path.exists(TEMPLATE_DOCX):
        raise FileNotFoundError(f"Template do Termo de Cautela não encontrado em: {TEMPLATE_DOCX}")

    doc = docx.Document(TEMPLATE_DOCX)

    numero = str(dados.get("numero") or "001").strip()
    ano = str(dados.get("ano") or datetime.now().year).strip()
    proc_sipac = str(dados.get("processo_sipac") or "").strip()
    vig_inicio = formatar_data_br(dados.get("vigencia_inicio"))
    vig_fim = formatar_data_br(dados.get("vigencia_fim"))
    data_extenso = formatar_data_extenso(dados.get("data_emissao"))

    direcao = dados.get("direcao") or {}
    dir_nome = str(direcao.get("nome") or "Autoridade Designada").strip()
    dir_cargo = str(direcao.get("cargo") or "Diretor Geral do Polo de Inovação Manaus").strip()
    dir_ato_tipo = str(direcao.get("ato_tipo") or "Portaria").strip()
    dir_ato_num = str(direcao.get("ato_numero") or "S/N").strip()
    dir_ato_origem = str(direcao.get("ato_origem") or "GR/IFAM").strip()
    dir_ato_data = str(direcao.get("ato_data") or "15/01/2024").strip()
    dir_ato_completo = f"{dir_ato_tipo} nº {dir_ato_num}/{dir_ato_origem}, de {dir_ato_data}"

    coord = dados.get("coordenacao") or {}
    coord_nome = str(coord.get("nome") or "Coordenador Responsável").strip()
    coord_locus = str(coord.get("locus") or "Laboratório / Setor Técnico").strip()
    coord_siape = str(coord.get("siape") or "").strip()
    coord_ato = str(coord.get("ato") or "").strip()

    solic = dados.get("solicitante") or {}
    solic_nome = str(solic.get("nome") or "").strip().upper()
    solic_cpf = str(solic.get("cpf") or "").strip()
    solic_vinculo = str(solic.get("vinculo") or "Servidor do IFAM").strip()
    solic_cargo = str(solic.get("cargo") or "").strip()
    solic_funcao = str(solic.get("funcao") or "Pesquisador / Colaborador Técnico").strip()
    solic_siape = str(solic.get("siape") or "").strip()
    solic_origem = str(solic.get("origem") or "Polo de Inovação Manaus").strip()
    solic_projeto = str(solic.get("projeto") or "Atividades do Polo de Inovação Manaus").strip()

    is_servidor = "servidor" in solic_vinculo.lower()

    if is_servidor:
        str_cargo = solic_cargo or "Servidor do IFAM"
        str_siape = solic_siape if solic_siape else "Em cadastramento"
    else:
        str_cargo = solic_cargo or solic_vinculo or "Colaborador Externo"
        str_siape = "Não aplicável (Colaborador Externo)"

    # 1. PARÁGRAFOS DO DOCUMENTO
    for p in doc.paragraphs:
        txt = p.text
        if not ("{{" in txt or "{%" in txt):
            continue

        if "cautela.numero" in txt:
            substituir_texto_preservando_estilo(p, f"Nº {numero}/{ano}")
            continue

        if "processo.numero" in txt:
            sipac_txt = proc_sipac if proc_sipac else "____________________________________"
            substituir_texto_preservando_estilo(p, f"Processo Administrativo SIPAC nº: {sipac_txt}")
            continue

        if "UNIDADE OUTORGANTE" in txt:
            novo_p4 = (
                f"1.1. UNIDADE OUTORGANTE: POLO DE INOVAÇÃO MANAUS / IFAM, "
                f"neste ato representado por seu {dir_cargo}, {dir_nome}, "
                f"nomeado pelo(a) {dir_ato_tipo} nº {dir_ato_num}/{dir_ato_origem}, de {dir_ato_data}."
            )
            substituir_texto_preservando_estilo(p, novo_p4)
            continue

        if "COORDENAÇÃO DE VINCULAÇÃO" in txt:
            siape_txt = f"Matrícula SIAPE nº {coord_siape}, " if coord_siape else ""
            novo_p5 = (
                f"1.2. COORDENAÇÃO DE VINCULAÇÃO: {coord_nome}, {coord_locus}, "
                f"{siape_txt}responsável pela gestão técnica do {coord_locus}."
            )
            substituir_texto_preservando_estilo(p, novo_p5)
            continue

        if "Nome Completo:" in txt:
            substituir_texto_preservando_estilo(p, f"Nome Completo: {solic_nome}")
            continue

        if "Cargo:" in txt:
            substituir_texto_preservando_estilo(p, f"Cargo: {str_cargo}")
            continue

        if "Função no Projeto:" in txt or "Funcao no Projeto:" in txt:
            substituir_texto_preservando_estilo(p, f"Função no Projeto: {solic_funcao}")
            continue

        if "SIAPE:" in txt and ("solicitante.pessoa.vinculo" in txt or "pessoa.siape" in txt):
            substituir_texto_preservando_estilo(p, f"SIAPE: {str_siape}")
            continue

        if "CPF:" in txt and ("solicitante.pessoa.vinculo" in txt or "pessoa.cpf" in txt):
            substituir_texto_preservando_estilo(p, f"CPF: {solic_cpf}")
            continue

        if "Lotação / Unidade de Origem:" in txt or "Lotacao" in txt:
            substituir_texto_preservando_estilo(p, f"Lotação / Unidade de Origem: {solic_origem}")
            continue

        if "solicitante.projeto.nome" in txt:
            novo_p16 = (
                f"3.1. O(s) equipamento(s) destina(m)-se exclusivamente à execução das atividades "
                f"técnicas, científicas e operacionais vinculadas ao projeto {solic_projeto}, "
                f"sendo expressamente vedada sua utilização para fins particulares ou em atividades "
                f"alheias às competências autorizadas pela Direção do Polo de Inovação Manaus."
            )
            substituir_texto_preservando_estilo(p, novo_p16)
            continue

        if "cautela.vigencia.inicio" in txt:
            novo_p18 = (
                f"4.1. A presente cautela vigorará de {vig_inicio} até {vig_fim}, "
                f"adstrita ao cronograma de execução das atividades do respectivo projeto."
            )
            substituir_texto_preservando_estilo(p, novo_p18)
            continue

        if "Manaus (AM)" in txt and "cautela.data" in txt:
            substituir_texto_preservando_estilo(p, f"Manaus (AM), {data_extenso}.")
            continue

        # Assinatura Direção
        if "direcao.pessoa.nome" in txt or "direção.pessoa.nome" in txt:
            substituir_texto_preservando_estilo(p, dir_nome)
            continue
        if "direção.nome" in txt or "direcao.nome" in txt:
            substituir_texto_preservando_estilo(p, dir_cargo)
            continue
        if "ato_nomeacao.tipo.numero.origem.data" in txt:
            substituir_texto_preservando_estilo(p, dir_ato_completo)
            continue

        # Assinatura Coordenação
        if "coordenacao.pessoa.nome" in txt or "coordenação.pessoa.nome" in txt:
            substituir_texto_preservando_estilo(p, coord_nome)
            continue
        if "coordenação.nome" in txt or "coordenacao.nome" in txt:
            substituir_texto_preservando_estilo(p, coord_locus)
            continue

        # Assinatura Cautelado
        if "solicitante.pessoa.nome" in txt:
            substituir_texto_preservando_estilo(p, solic_nome)
            continue
        if "solicitante.pessoa.siape" in txt:
            if is_servidor and solic_siape:
                substituir_texto_preservando_estilo(p, f"SIAPE nº {solic_siape}")
            else:
                substituir_texto_preservando_estilo(p, "")
            continue
        if "solicitante.pessoa.cpf" in txt:
            substituir_texto_preservando_estilo(p, f"CPF nº {solic_cpf}")
            continue

    # 2. TABELA DE EQUIPAMENTOS (TABELA 0)
    if doc.tables:
        tabela = doc.tables[0]
        while len(tabela.rows) > 1:
            tr = tabela.rows[-1]._tr
            tabela._tbl.remove(tr)

        itens = dados.get("itens") or []
        if not itens:
            row = tabela.add_row()
            row.cells[0].text = "1"
            row.cells[1].text = "EQUIPAMENTO CONFORME ESPECIFICADO EM TERMO ADITIVO"
            row.cells[2].text = "S/N"
            row.cells[3].text = "Bom estado"
            row.cells[4].text = "-"
        else:
            for idx, item in enumerate(itens, start=1):
                row = tabela.add_row()
                desc = str(item.get("descricao") or item.get("nome_material") or "").strip()
                serie = str(item.get("numero_serie") or item.get("patrimonio") or "S/N").strip()
                estado = str(item.get("estado_conservacao") or "Bom estado de conservação e funcionamento").strip()
                acess = str(item.get("acessorios") or item.get("observacao") or "-").strip()

                row.cells[0].text = str(idx)
                row.cells[1].text = desc
                row.cells[2].text = serie
                row.cells[3].text = estado
                row.cells[4].text = acess

                for ci, cell in enumerate(row.cells):
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    if cell.paragraphs:
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(3)
                        p.paragraph_format.space_after = Pt(3)
                        if ci == 0:
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        elif ci in (2, 3):
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        else:
                            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        for r in p.runs:
                            r.font.name = "Calibri"
                            r.font.size = Pt(9.5)

    nome_sanitizado = sanitizar_nome_arquivo(solic_nome if solic_nome else "Cautelado")
    nome_arquivo = f"Termo_Cautela_{numero}_{ano}_{nome_sanitizado}.docx"

    if not caminho_saida:
        caminho_saida = os.path.join(DIR_SAIDA_CAUTELAS, nome_arquivo)

    doc.save(caminho_saida)

    return {
        "sucesso": True,
        "caminho_arquivo": caminho_saida,
        "nome_arquivo": nome_arquivo,
        "numero": numero,
        "ano": ano,
        "codigo_cautela": f"{numero}/{ano}"
    }