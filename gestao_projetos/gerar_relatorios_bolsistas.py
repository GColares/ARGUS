"""
Script para geração automatizada de Relatórios de Atividades em formato .docx
a partir do modelo Word e da planilha Excel de dados dos bolsistas.
Atualizado para usar a biblioteca docxtpl (Jinja2) garantindo a integridade
dos estilos, formatações e XML nativo do Word sem corromper o arquivo.
"""

import os
import datetime
import openpyxl
from docxtpl import DocxTemplate

def gerar_relatorios(excel_path=None, template_path=None, output_dir=None):
    base_dir = os.path.dirname(os.path.abspath(__file__))
    modelos_dir = os.path.join(base_dir, 'media', 'modelos')
    
    if not output_dir:
        output_dir = modelos_dir
        
    if not excel_path:
        excel_path = os.path.join(modelos_dir, 'parcela1-b-2.xlsx')
        
    if not template_path:
        # Encontra o modelo .docx na pasta modelos (ignora arquivos temporários ~$)
        templates = [
            f for f in os.listdir(modelos_dir)
            if f.endswith('.docx') and not f.startswith('~$') and ('Rascunho' in f or 'Modelo' in f)
        ]
        if not templates:
            raise FileNotFoundError("Nenhum arquivo de modelo (.docx) encontrado na pasta de modelos.")
        template_path = os.path.join(modelos_dir, templates[0])

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb.active

    # Extrai cabeçalhos da primeira linha
    headers = []
    for cell in next(ws.iter_rows(values_only=True)):
        if cell is not None and str(cell).strip():
            headers.append(str(cell).strip())
        else:
            break

    today_str = datetime.date.today().strftime('%d/%m/%Y')
    relatorios_gerados = []

    for r in ws.iter_rows(min_row=2, values_only=True):
        # Ignora linhas totalmente vazias ou sem os primeiros dados chave
        if not r or all(cell is None or str(cell).strip() == '' for cell in r[:5]):
            continue
            
        data_dict = {}
        for i in range(len(headers)):
            val = r[i] if i < len(r) and r[i] is not None else ''
            data_dict[headers[i]] = str(val).strip()
            
        data_dict['hoje'] = today_str
        
        # Mapeamento de Aliases e Novas Tags
        data_dict['bolsista_rg_numero'] = data_dict.get('bolsista_rg_numero') or data_dict.get('bolsista_rg', '')
        data_dict['bolsista_termodebolsa_numero'] = data_dict.get('bolsista_termodebolsa_numero') or data_dict.get('bolsista_termodebolsa', '')
        
        data_dict['parcela_numero'] = data_dict.get('parcela_numero') or data_dict.get('parcela-n', '') or data_dict.get('n-parcela', '')
        data_dict['parcela_total'] = data_dict.get('parcela_total') or data_dict.get('parcela-total', '') or data_dict.get('total-parcela', '')
        data_dict['parcela_ch'] = data_dict.get('parcela_ch') or data_dict.get('parcela_CH', '')
        
        if 'contratação_modalidade' not in data_dict and 'contratacao_modalidade' not in data_dict:
            data_dict['contratação_modalidade'] = 'Bolsa'

        bolsista_nome = data_dict.get('bolsista_nome', 'Bolsista').strip()
        p_n = data_dict.get('parcela_numero', '1').strip()
        p_tot = data_dict.get('parcela_total', '1').strip()
        
        if not bolsista_nome:
            continue

        out_filename = f'Relatorio_de_Atividades_{bolsista_nome}_-_Parcela_{p_n}_de_{p_tot}.docx'
        out_path = os.path.join(output_dir, out_filename)
        
        # Processamento seguro com docxtpl preservando todo o XML/Word
        doc = DocxTemplate(template_path)
        doc.render(data_dict)
        doc.save(out_path)
                        
        relatorios_gerados.append(out_path)
        print(f"Gerado com sucesso (docxtpl): {out_filename}")

    return relatorios_gerados

if __name__ == '__main__':
    gerados = gerar_relatorios()
    print(f"\nTotal de relatórios gerados com integridade garantida: {len(gerados)}")
